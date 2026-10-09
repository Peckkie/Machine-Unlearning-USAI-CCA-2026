"""
Full metric table for the paper (X1 / X2 + A1 / A3 / A5), written to one Excel workbook.

Reads every per-image prediction file written by evaluate.py ($PRED_DIR/{model}_seed{k}_{lab|field}.csv),
so it can simply be re-run when C1 / C2 seeds finish.

Per model x seed x test set:
  accuracy (+ Clopper-Pearson 95% CI), balanced accuracy, macro precision / recall / F1
  (Table 3/4 convention: labels = classes in y_true U y_pred), macro F1 over classes present in y_true,
  weighted F1, Cohen's kappa, top-2 / top-3 accuracy, macro one-vs-rest ROC-AUC (classes present in y_true),
  Normal-vs-Abnormal sensitivity / specificity / PPV / NPV / F1 / ROC-AUC (score = 1 - p_Normal),
  case-level bootstrap 95% CI for the main metrics.
Per class: precision, sensitivity, specificity, F1, ROC-AUC, support (mean +- SD over seeds).
Comparisons vs original (same seed): delta with case-level bootstrap CI + exact McNemar; pooled over the
seeds both models have (same resampled patients for every pair).

usage (local, after copying predictions + test CSVs):
  python full_metrics.py --pred_dir pred --lab_csv Testdf_fold1_2_v1.csv \
      --field_csv UICCA_DiagRadioExp_AzureDb51Case813Image.csv --analysis_dir analysis --out Paper_Metrics.xlsx
"""
import argparse
import glob
import os
import re

import numpy as np
import pandas as pd
from scipy.stats import beta
from sklearn.metrics import cohen_kappa_score, confusion_matrix, roc_auc_score

try:
    from scipy.stats import binomtest
    def binom_p(k, n): return binomtest(k, n, 0.5).pvalue
except ImportError:  # scipy < 1.7
    from scipy.stats import binom_test
    def binom_p(k, n): return binom_test(k, n, 0.5)

MODEL_ORDER = ['original', 'unlearned', 'C1', 'C2']
MODEL_NAME = {'original': 'Original (ImageNet init)', 'unlearned': 'Unlearned (proposed)',
              'C1': 'C1 compute-matched fine-tuning', 'C2': 'C2 same/different-image auxiliary'}
B, RNG_SEED = 1000, 2026
CI_METRICS = ['accuracy', 'balanced_accuracy', 'macro_f1', 'abn_sensitivity', 'abn_specificity', 'abn_auc']
CMP_METRICS = ['accuracy', 'balanced_accuracy', 'macro_f1', 'macro_precision', 'macro_recall',
               'abn_sensitivity', 'abn_specificity', 'abn_auc']


def cp_ci(k, n, a=0.05):
    lo = beta.ppf(a / 2, k, n - k + 1) if k > 0 else 0.0
    hi = beta.ppf(1 - a / 2, k + 1, n - k) if k < n else 1.0
    return lo, hi


def core_metrics(y, p, pab):
    """fast metrics used inside the bootstrap; y/p are int-coded, pab = P(abnormal)."""
    labs = np.union1d(y, p)
    tp = np.array([((y == c) & (p == c)).sum() for c in labs], float)
    npred = np.array([(p == c).sum() for c in labs], float)
    ntrue = np.array([(y == c).sum() for c in labs], float)
    prec = np.divide(tp, npred, out=np.zeros_like(tp), where=npred > 0)
    rec = np.divide(tp, ntrue, out=np.zeros_like(tp), where=ntrue > 0)
    f1 = np.divide(2 * prec * rec, prec + rec, out=np.zeros_like(tp), where=(prec + rec) > 0)
    present = ntrue > 0
    ay, ap = y != NORMAL, p != NORMAL
    try:
        auc = roc_auc_score(ay, pab) if 0 < ay.sum() < len(ay) else np.nan
    except ValueError:
        auc = np.nan
    return dict(accuracy=(y == p).mean(), balanced_accuracy=rec[present].mean(),
                macro_precision=prec.mean(), macro_recall=rec.mean(), macro_f1=f1.mean(),
                abn_sensitivity=(ay & ap).sum() / max(ay.sum(), 1),
                abn_specificity=(~ay & ~ap).sum() / max((~ay).sum(), 1), abn_auc=auc)


def full_metrics(y, p, P, classes):
    m = core_metrics(y, p, 1 - P[:, NORMAL])
    n, k = len(y), int((y == p).sum())
    m['accuracy_ci_cp'] = '%.3f-%.3f' % cp_ci(k, n)
    present = np.unique(y)
    labs = np.union1d(y, p)
    f1s = []
    for c in present:
        tp = ((y == c) & (p == c)).sum()
        pr = tp / max((p == c).sum(), 1)
        rc = tp / max((y == c).sum(), 1)
        f1s.append(2 * pr * rc / (pr + rc) if pr + rc else 0.0)
    m['macro_f1_present_classes'] = np.mean(f1s)
    w = np.array([(y == c).sum() for c in present], float)
    m['weighted_f1'] = np.sum(np.array(f1s) * w) / w.sum()
    m['cohen_kappa'] = cohen_kappa_score(y, p)
    order = np.argsort(-P, axis=1)
    m['top2_accuracy'] = np.mean([y[i] in order[i, :2] for i in range(n)])
    m['top3_accuracy'] = np.mean([y[i] in order[i, :3] for i in range(n)])
    aucs = [roc_auc_score(y == c, P[:, c]) for c in present if 0 < (y == c).sum() < n]
    m['macro_auc_ovr'] = np.mean(aucs)
    ay, ap = y != NORMAL, p != NORMAL
    m['abn_ppv'] = (ay & ap).sum() / max(ap.sum(), 1)
    m['abn_npv'] = (~ay & ~ap).sum() / max((~ap).sum(), 1)
    s, pv = m['abn_sensitivity'], m['abn_ppv']
    m['abn_f1'] = 2 * s * pv / (s + pv) if s + pv else 0.0
    m['n_images'], m['n_classes_true'], m['n_classes_pred'] = n, len(present), len(np.unique(p))
    return m


def load(pred_dir, lab_csv, field_csv):
    refs = {'lab': (lab_csv, 'Path Crop', lambda d: d[(d['Path Crop'] != 'None') & (d['Path Crop'] != 'Nan')]),
            'field': (field_csv, 'Path Crop_I7', lambda d: d[d['Sub_Class_15AB'] != 'None'])}
    data = {}
    for ts, (csv, pcol, keep) in refs.items():
        ref = keep(pd.read_csv(csv, keep_default_na=False)).reset_index(drop=True)  # newer pandas reads 'None' as NaN
        runs = {}
        for f in sorted(glob.glob(f'{pred_dir}/*_seed*_{ts}.csv')):
            mm = re.match(r'(.+)_seed(\d+)_' + ts + r'\.csv', os.path.basename(f))
            d = pd.read_csv(f)
            assert len(d) == len(ref) and (d['img_path'].values == ref[pcol].values).all(), f'{f}: row order mismatch'
            runs[(mm.group(1), int(mm.group(2)))] = d
        data[ts] = (ref['Case'].values, runs)
    return data


def main():
    global NORMAL
    ap = argparse.ArgumentParser()
    ap.add_argument('--pred_dir', required=True)
    ap.add_argument('--lab_csv', required=True)
    ap.add_argument('--field_csv', required=True)
    ap.add_argument('--analysis_dir', default=None, help='folder with A1_face_verification.csv / A3_*.csv')
    ap.add_argument('--out', default='Paper_Metrics.xlsx')
    args = ap.parse_args()
    rng = np.random.default_rng(RNG_SEED)
    data = load(args.pred_dir, args.lab_csv, args.field_csv)

    per_run, per_class, cms, pairs, pooled = [], [], [], [], []
    for ts, (cases, runs) in data.items():
        any_df = next(iter(runs.values()))
        classes = [c[2:] for c in any_df.columns if c.startswith('p_')]
        cidx = {c: i for i, c in enumerate(classes)}
        NORMAL = cidx['Normal']
        y = any_df['label'].map(cidx).values
        R = {k: (d['pred'].map(cidx).values, d[[f'p_{c}' for c in classes]].values) for k, d in runs.items()}
        uniq = np.unique(cases)
        idx_by_case = {c: np.where(cases == c)[0] for c in uniq}
        boots = [np.concatenate([idx_by_case[c] for c in rng.choice(uniq, len(uniq))]) for _ in range(B)]

        for (model, seed), (p, P) in sorted(R.items(), key=lambda kv: (MODEL_ORDER.index(kv[0][0]), kv[0][1])):
            m = full_metrics(y, p, P, classes)
            bt = [core_metrics(y[i], p[i], 1 - P[i, NORMAL]) for i in boots]
            for mt in CI_METRICS:
                a = np.array([b[mt] for b in bt], float)
                m[f'{mt}_ci_case'] = '%.3f-%.3f' % (np.nanpercentile(a, 2.5), np.nanpercentile(a, 97.5))
            per_run.append(dict(testset=ts, model=model, seed=seed, n_cases=len(uniq), **m))
            cm = confusion_matrix(y, p, labels=range(len(classes)))
            for i, ci in enumerate(classes):
                cms.append(dict(testset=ts, model=model, seed=seed, true=ci, **{c: int(cm[i, j]) for j, c in enumerate(classes)}))
            for i, c in enumerate(classes):
                sup = int((y == i).sum())
                if sup == 0:
                    continue
                tp = int(cm[i, i]); fp = int(cm[:, i].sum() - tp); fn = sup - tp; tn = len(y) - tp - fp - fn
                pr = tp / (tp + fp) if tp + fp else 0.0; rc = tp / sup
                auc = roc_auc_score(y == i, P[:, i]) if sup < len(y) else np.nan
                per_class.append(dict(testset=ts, model=model, seed=seed, cls=c, support=sup, precision=pr,
                                      sensitivity=rc, specificity=tn / (tn + fp), f1=2 * pr * rc / (pr + rc) if pr + rc else 0.0,
                                      auc=auc))

        # comparisons vs original
        for model in MODEL_ORDER[1:]:
            seeds = sorted(s for (mm, s) in R if mm == model and ('original', s) in R)
            if not seeds:
                continue
            for s in seeds:
                (po, Po), (pu, Pu) = R[('original', s)], R[(model, s)]
                mo, mu = core_metrics(y, po, 1 - Po[:, NORMAL]), core_metrics(y, pu, 1 - Pu[:, NORMAL])
                bd = [{mt: core_metrics(y[i], pu[i], 1 - Pu[i, NORMAL])[mt] - core_metrics(y[i], po[i], 1 - Po[i, NORMAL])[mt]
                       for mt in CMP_METRICS} for i in boots]
                co, cu = po == y, pu == y
                b_, c_ = int((co & ~cu).sum()), int((~co & cu).sum())
                row = dict(testset=ts, comparison=f'{model} vs original', seed=s,
                           mcnemar_only_original_correct=b_, mcnemar_only_model_correct=c_,
                           mcnemar_exact_p=binom_p(min(b_, c_), b_ + c_) if b_ + c_ else 1.0)
                for mt in CMP_METRICS:
                    a = np.array([d[mt] for d in bd], float)
                    row.update({f'{mt}_original': mo[mt], f'{mt}_model': mu[mt], f'{mt}_delta': mu[mt] - mo[mt],
                                f'{mt}_delta_ci': '%+.3f to %+.3f' % (np.nanpercentile(a, 2.5), np.nanpercentile(a, 97.5))})
                pairs.append(row)
            bd = []
            for i in boots:
                ds = []
                for s in seeds:
                    (po, Po), (pu, Pu) = R[('original', s)], R[(model, s)]
                    a_, b2 = core_metrics(y[i], pu[i], 1 - Pu[i, NORMAL]), core_metrics(y[i], po[i], 1 - Po[i, NORMAL])
                    ds.append({mt: a_[mt] - b2[mt] for mt in CMP_METRICS})
                bd.append({mt: np.mean([d[mt] for d in ds]) for mt in CMP_METRICS})
            for mt in CMP_METRICS:
                o = [core_metrics(y, *(lambda t: (t[0], 1 - t[1][:, NORMAL]))(R[('original', s)]))[mt] for s in seeds]
                u = [core_metrics(y, *(lambda t: (t[0], 1 - t[1][:, NORMAL]))(R[(model, s)]))[mt] for s in seeds]
                a = np.array([d[mt] for d in bd], float)
                pooled.append(dict(testset=ts, comparison=f'{model} vs original', seeds=','.join(map(str, seeds)),
                                   metric=mt, original_mean=np.mean(o), original_sd=np.std(o, ddof=1) if len(o) > 1 else np.nan,
                                   model_mean=np.mean(u), model_sd=np.std(u, ddof=1) if len(u) > 1 else np.nan,
                                   mean_delta=np.mean(u) - np.mean(o), ci_lo=np.nanpercentile(a, 2.5), ci_hi=np.nanpercentile(a, 97.5),
                                   seeds_model_better=f'{sum(ui > oi for oi, ui in zip(o, u))}/{len(seeds)}'))

    pr = pd.DataFrame(per_run)
    pc = pd.DataFrame(per_class)
    num = ['accuracy', 'balanced_accuracy', 'macro_precision', 'macro_recall', 'macro_f1', 'macro_f1_present_classes',
           'weighted_f1', 'cohen_kappa', 'top2_accuracy', 'top3_accuracy', 'macro_auc_ovr',
           'abn_sensitivity', 'abn_specificity', 'abn_ppv', 'abn_npv', 'abn_f1', 'abn_auc']

    # summary: mean +- SD over seeds (string for the paper + numeric columns)
    summ = []
    for (ts, model), g in pr.groupby(['testset', 'model'], sort=False):
        row = dict(testset=ts, model=MODEL_NAME[model], seeds_done=len(g))
        for mt in num:
            v = g[mt].values
            row[mt] = f'{v.mean():.3f} ± {v.std(ddof=1):.3f}' if len(v) > 1 else f'{v.mean():.3f} (1 seed)'
        summ.append(row)
    summ = pd.DataFrame(summ)
    pcs = (pc.groupby(['testset', 'cls', 'model'], sort=False)
             .agg(support=('support', 'first'), n_seeds=('seed', 'count'),
                  precision_mean=('precision', 'mean'), precision_sd=('precision', 'std'),
                  sensitivity_mean=('sensitivity', 'mean'), sensitivity_sd=('sensitivity', 'std'),
                  specificity_mean=('specificity', 'mean'), specificity_sd=('specificity', 'std'),
                  f1_mean=('f1', 'mean'), f1_sd=('f1', 'std'), auc_mean=('auc', 'mean'), auc_sd=('auc', 'std'))
             .reset_index())

    write_excel(args, summ, pr, pd.DataFrame(pairs), pd.DataFrame(pooled), pcs, pc, pd.DataFrame(cms))
    print(summ.to_string())
    print('saved', args.out)


def write_excel(args, summ, pr, pairs, pooled, pcs, pc, cms):
    status = pr.groupby('model').seed.apply(lambda s: ','.join(map(str, sorted(set(s))))).to_dict()
    readme = pd.DataFrame({'item': [
        'Source', 'Test sets', 'Seeds available', 'Macro P/R/F1', 'Macro F1 (present classes)', 'Balanced accuracy',
        'Abnormal (abn_*)', 'abn_auc', 'macro_auc_ovr', 'top2/top3', 'CI (case)', 'CI (CP)', 'Comparisons', 'McNemar',
        'Pooled', 'A1', 'A3', 'A5'],
        'description': [
        'per-image predictions from experiments/evaluate.py; re-run experiments/analysis/full_metrics.py when C1/C2 seeds finish',
        'lab = Testdf_fold1_2_v1 (1,312 images, 144 cases); field = UICCA Azure 51 cases (807 images, Sub_Class_15AB != None)',
        '; '.join(f'{MODEL_NAME[k]}: seeds {v}' for k, v in sorted(status.items(), key=lambda kv: MODEL_ORDER.index(kv[0]))),
        'sklearn classification_report macro avg, labels = classes in y_true U y_pred (same as Table 3/4 of the paper)',
        'macro F1 averaged only over classes that occur in the test set (classes predicted but absent are excluded)',
        'mean recall over classes present in the test set',
        'binary Normal vs any abnormal (AB01-AB12): sensitivity, specificity, PPV, NPV, F1',
        'ROC-AUC of Normal vs Abnormal with score = 1 - p(Normal)',
        'one-vs-rest ROC-AUC averaged over classes present in the test set',
        'true class among the 2 / 3 highest softmax outputs',
        f'95% percentile CI from case-level (patient) cluster bootstrap, B = {B}, seed {RNG_SEED}',
        'exact Clopper-Pearson 95% CI of image-level accuracy',
        'model seed k vs original seed k (same pipeline, same seed); delta = model - original, CI from case bootstrap',
        'exact two-sided McNemar on per-image correctness (binomial test on discordant pairs)',
        'mean delta over seeds; bootstrap resamples the same patients for every seed pair',
        'face verification (30 identities): ROC-AUC, EER, TAR@FAR, identity-level bootstrap CI',
        'flip consistency cos(f(x), f(flip x)) on backbone / downstream features; Wilcoxon signed-rank',
        'GPU-hours from TensorBoard wall-clock (unlearning on RTX 3090 Ti, downstream on RTX 2080 Ti)']})

    sheets = [('README', readme), ('Summary_mean_SD', summ), ('Per_seed', pr), ('Compare_pooled', pooled),
              ('Compare_per_seed', pairs), ('Per_class_mean_SD', pcs), ('Per_class_per_seed', pc), ('Confusion', cms)]
    if args.analysis_dir:
        f = f'{args.analysis_dir}/A1_face_verification.csv'
        if os.path.exists(f):
            sheets.append(('A1_face_verification', pd.read_csv(f)))
        a3 = []
        for ts in ('lab', 'field'):
            for kind in ('summary', 'pairs'):
                f = f'{args.analysis_dir}/A3_flip_consistency_{ts}_{kind}.csv'
                if os.path.exists(f):
                    d = pd.read_csv(f); d.insert(0, 'table', kind); d.insert(0, 'testset', ts); a3.append(d)
        if a3:
            sheets.append(('A3_flip_consistency', pd.concat(a3, ignore_index=True)))
    sheets.append(('A5_compute', pd.DataFrame([
        dict(stage='Unlearning R1 (flip task, mini-ImageNet, 200 ep)', gpu='RTX 3090 Ti', gpu_hours=187.4),
        dict(stage='Unlearning R2 (unfreeze B4-B7, 50 ep)', gpu='RTX 3090 Ti', gpu_hours=45.5),
        dict(stage='Unlearning total', gpu='RTX 3090 Ti', gpu_hours=232.9),
        dict(stage='Downstream fine-tuning per model (R1 200 + R2 200 ep)', gpu='RTX 2080 Ti', gpu_hours=28.0)])))

    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(args.out, engine='openpyxl') as xw:
        for name, df in sheets:
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = 'A2'
            for c in ws[1]:
                c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='305496')
                c.alignment = Alignment(wrap_text=True, vertical='center')
            for j, col in enumerate(df.columns, 1):
                w = max([len(str(col))] + [len(f'{v:.4f}' if isinstance(v, float) else str(v)) for v in df[col].head(200)])
                ws.column_dimensions[get_column_letter(j)].width = min(max(w + 2, 8), 80 if name == 'README' else 30)
                if df[col].dtype.kind == 'f':
                    fmt = '0.000E+00' if 'p' == col.split('_')[-1] or col.endswith('_p') else '0.0000'
                    for r in range(2, len(df) + 2):
                        ws.cell(r, j).number_format = fmt


if __name__ == '__main__':
    main()
